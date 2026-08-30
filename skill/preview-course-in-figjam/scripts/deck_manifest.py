#!/usr/bin/env python3
"""Extract a token-efficient slide manifest from PDF or PPTX input."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import re
import sys
import tempfile
import unicodedata


SCHEMA_VERSION = 1
LOW_INFORMATION_LIMIT = 40


def normalize_text(value: str) -> str:
    value = unicodedata.normalize("NFC", value or "").replace("\x00", "")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in value.splitlines()]
    compact: list[str] = []
    blank = False
    for line in lines:
        if line:
            compact.append(line)
            blank = False
        elif compact and not blank:
            compact.append("")
            blank = True
    return "\n".join(compact).strip()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def classify(text: str, image_count: int) -> str:
    count = len(text.strip())
    if count == 0:
        return "image_only" if image_count > 0 else "blank"
    if count < LOW_INFORMATION_LIMIT:
        return "low_information"
    return "content"


def inspect_pdf(path: Path) -> list[dict]:
    try:
        import pdfplumber
    except ImportError as exc:
        raise RuntimeError("pdfplumber is required; use the bundled workspace Python runtime") from exc

    slides: list[dict] = []
    with pdfplumber.open(path) as pdf:
        for index, page in enumerate(pdf.pages, start=1):
            text = normalize_text(page.extract_text() or "")
            images = len(page.images)
            slides.append(
                {
                    "pageNumber": index,
                    "text": text,
                    "charCount": len(text),
                    "imageCount": images,
                    "widthPoints": round(float(page.width), 2),
                    "heightPoints": round(float(page.height), 2),
                    "classification": classify(text, images),
                }
            )
    return slides


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if getattr(shape, "shape_type", None) == 6 and hasattr(shape, "shapes"):
            yield from iter_shapes(shape.shapes)


def inspect_pptx(path: Path) -> list[dict]:
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as exc:
        raise RuntimeError("python-pptx is required; use the bundled workspace Python runtime") from exc

    presentation = Presentation(str(path))
    width_points = round(presentation.slide_width / 12700, 2)
    height_points = round(presentation.slide_height / 12700, 2)
    slides: list[dict] = []

    for index, slide in enumerate(presentation.slides, start=1):
        text_parts: list[str] = []
        image_count = 0
        for shape in iter_shapes(slide.shapes):
            if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
                image_count += 1
            if getattr(shape, "has_text_frame", False):
                value = normalize_text(getattr(shape, "text", ""))
                if value:
                    text_parts.append(value)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        value = normalize_text(cell.text)
                        if value:
                            text_parts.append(value)
        text = normalize_text("\n".join(text_parts))
        slides.append(
            {
                "pageNumber": index,
                "text": text,
                "charCount": len(text),
                "imageCount": image_count,
                "widthPoints": width_points,
                "heightPoints": height_points,
                "classification": classify(text, image_count),
            }
        )
    return slides


def write_json_atomic(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as handle:
            json.dump(payload, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def build_manifest(source: Path) -> dict:
    suffix = source.suffix.casefold()
    if suffix == ".pdf":
        slides = inspect_pdf(source)
        source_format = "pdf"
    elif suffix == ".pptx":
        slides = inspect_pptx(source)
        source_format = "pptx"
    else:
        raise ValueError("SOURCE must be a .pdf or .pptx file")

    counts = {kind: 0 for kind in ("content", "low_information", "image_only", "blank")}
    for slide in slides:
        counts[slide["classification"]] += 1
    return {
        "schemaVersion": SCHEMA_VERSION,
        "source": {
            "name": source.name,
            "format": source_format,
            "sha256": sha256_file(source),
        },
        "slideCount": len(slides),
        "classificationCounts": counts,
        "slides": slides,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()

    source = args.source.expanduser().resolve()
    if not source.is_file():
        parser.error(f"source file does not exist: {source}")
    try:
        manifest = build_manifest(source)
        if args.output:
            write_json_atomic(args.output.expanduser().resolve(), manifest)
        else:
            json.dump(manifest, sys.stdout, ensure_ascii=False, indent=2)
            sys.stdout.write("\n")
    except Exception as exc:
        print(f"deck_manifest: {exc}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
